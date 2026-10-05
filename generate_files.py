import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import pptx
from pptx import Presentation
from pptx.util import Inches as PPTInches, Pt as PPTPt
from pptx.dml.color import RGBColor as PPTRGBColor
from pptx.enum.text import PP_ALIGN

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def generate_word_document():
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style definitions
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Title Banner / Header
    p_institution = doc.add_paragraph()
    p_institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_institution.add_run("CENTRO UNIVERSITÁRIO UNA / ÂNIMA EDUCAÇÃO\n")
    run_inst.bold = True
    run_inst.font.size = Pt(14)
    run_inst.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    run_sub = p_institution.add_run("RELATÓRIO TÉCNICO FINAL DE A3 — GARANTIA E QUALIDADE DE SOFTWARE")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    doc.add_paragraph() # Spacing

    # Metadata Panel
    table_meta = doc.add_table(rows=2, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    meta_data = [
        [("Disciplina:", "Garantia e Qualidade de Software (GQS)"), ("Professor:", "Daniel Henrique Matos de Paiva")],
        [("Data de Entrega:", "16/10/2026"), ("Projeto:", "EduGrade - Gestão & Qualidade de Software (ODS 4)")]
    ]

    for row_idx, row in enumerate(meta_data):
        for col_idx, (label, val) in enumerate(row):
            cell = table_meta.cell(row_idx, col_idx)
            set_cell_background(cell, "F7FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run(f"{label} ")
            r1.bold = True
            r1.font.size = Pt(10)
            r2 = p.add_run(val)
            r2.font.size = Pt(10)

    doc.add_paragraph()

    # Section Helper
    def add_heading_1(text):
        h = doc.add_heading(text, level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    def add_heading_2(text):
        h = doc.add_heading(text, level=2)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    def add_bullet(p_text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        p.add_run(p_text)

    # 1. Integrantes
    add_heading_1("1. IDENTIFICAÇÃO DOS INTEGRANTES DO GRUPO")
    
    p_repo = doc.add_paragraph()
    r_repo_label = p_repo.add_run("Link do Repositório GitHub: ")
    r_repo_label.bold = True
    p_repo.add_run("https://github.com/vivoeasy100/calculadora_de_notas_academicas-")

    table_members = doc.add_table(rows=5, cols=4)
    table_members.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Nome Completo", "RA", "Usuário GitHub", "Função no Projeto"]
    for col_idx, text in enumerate(headers):
        cell = table_members.cell(0, col_idx)
        set_cell_background(cell, "1A365D")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    members = [
        ("Lucas Henrique Miranda", "325131396", "@LuchMiranda", "Product Owner & Requisitos do Sistema"),
        ("Caio Duraes", "325132875", "@caiovas28-dotcom", "Desenvolvedor Backend & Arquitetura"),
        ("Gabriel Ferreira", "325140970", "@1Gapril", "DevOps & Gestão de Configuração (GitFlow)"),
        ("Fernando Almeida", "326132695", "@vivoeasy100", "Engenheiro de Qualidade & Testes (QA/TDD)")
    ]

    for row_idx, m in enumerate(members, start=1):
        bg = "EDF2F7" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(m):
            cell = table_members.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph()

    # 2. Problema do Mundo Real
    add_heading_1("2. PROBLEMA A SER RESOLVIDO DO MUNDO REAL")
    p_prob1 = doc.add_paragraph(
        "No ambiente universitário e escolar, milhares de estudantes enfrentam dificuldades frequentes na gestão "
        "de seu desempenho acadêmico ao longo do semestre letivo. Com sistemas institucionais complexos que utilizam "
        "critérios de médias ponderadas distintas (pesos entre avaliações A1, A2 e A3), provas integradas e cálculo de nota "
        "necessária para exame final/recuperação, muitos alunos só tomam conhecimento de sua situação de risco quando já é "
        "tarde demais para recuperação."
    )
    p_prob1.paragraph_format.space_after = Pt(6)
    
    p_prob2 = doc.add_paragraph("Essa falta de previsibilidade e acompanhamento contínuo contribui diretamente para:")
    add_bullet(" Desmotivação e ansiedade acadêmica durante os períodos de provas;")
    add_bullet(" Reprovações que poderiam ter sido evitadas com planejamento antecipado e monitoramento;")
    add_bullet(" Índices crescentes de evasão universitária pela falta de suporte decisório simples.")

    # 3. Solução Proposta
    add_heading_1("3. SOLUÇÃO PROPOSTA: SISTEMA \"EDUGRADE\"")
    doc.add_paragraph(
        "O EduGrade é uma solução computacional modular projetada para fornecer aos alunos e orientadores "
        "pedagógicos uma ferramenta prática, transparente e confiável para simulação, cálculo e acompanhamento de metas acadêmicas."
    )
    
    add_heading_2("Funcionalidades Implementadas:")
    add_bullet(" Permite personalizar pesos de cada avaliação institucional (A1, A2, A3) e calcular a média exata.", "1. Cálculo de Média Ponderada e Aritmética:")
    add_bullet(" Enquadramento imediato em Aprovado, Recuperação / Exame Final ou Reprovado conforme nota de corte.", "2. Diagnóstico Automático de Situação:")
    add_bullet(" Indica com precisão decimal a nota mínima necessária na prova de exame final para obtenção de aprovação.", "3. Cálculo Preditivo para Exame:")
    add_bullet(" Armazenamento de dados e histórico de estudantes em arquivo de persistência em JSON.", "4. Persistência de Dados em Formato JSON:")

    # 4. Vínculo com ODS
    add_heading_1("4. VÍNCULO COM OS OBJETIVOS DE DESENVOLVIMENTO SUSTENTÁVEL (ODS)")
    doc.add_paragraph(
        "O projeto está vinculado diretamente ao ODS 4: Educação de Qualidade da Organização das Nações Unidas (ONU), especificamente:"
    )
    add_bullet(" Capacitação dos estudantes em gestão de autonomia de aprendizado e desenvolvimento de ferramentas técnicas de qualidade.", "Meta 4.4 - Habilidades Técnicas: ")
    add_bullet(" A ferramenta apoia a redução da evasão através da autoavaliação contínua e previsibilidade de desempenho acadêmico.", "Incentivo à Permanência Estudantil: ")

    # 5. Boas Práticas GQS
    add_heading_1("5. APLICAÇÃO DAS BOAS PRÁTICAS DE GESTÃO E QUALIDADE DE SOFTWARE (GQS)")
    add_bullet(" Nomenclatura semântica, Single Responsibility Principle (SRP) e conformidade rigorosa com PEP 8.", "1. Clean Code (Código Limpo):")
    add_bullet(" Classes de exceção especializadas (AcademicError, InvalidGradeError, InvalidWeightError) com validação rígida de limites.", "2. Tratamento de Exceções Robusto:")
    add_bullet(" 10 testes unitários automatizados desenvolvidos com unittest, garantindo 100% de cobertura das regras de negócio.", "3. Desenvolvimento Guiado por Testes (TDD):")
    add_bullet(" Estrutura GitFlow (main, develop, feature/*) e padrão semântico de commits (feat, fix, test, docs).", "4. Versionamento e GitFlow:")
    add_bullet(" Pipeline automatizado no GitHub Actions (.github/workflows/ci.yml) executando validações em cada integração.", "5. Automação de CI/CD:")

    doc.save("Relatorio_Final_A3.docx")
    print("Relatorio_Final_A3.docx criado com sucesso!")

def generate_pptx_presentation():
    prs = Presentation()
    prs.slide_width = PPTInches(13.333)
    prs.slide_height = PPTInches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color scheme
    BG_DARK = PPTRGBColor(0x1A, 0x20, 0x2C)
    TEXT_WHITE = PPTRGBColor(0xFF, 0xFF, 0xFF)
    TEXT_MUTED = PPTRGBColor(0xA0, 0xAE, 0xC0)
    ACCENT_BLUE = PPTRGBColor(0x31, 0x82, 0xCE)
    ACCENT_GREEN = PPTRGBColor(0x38, 0xA1, 0x69)
    CARD_BG = PPTRGBColor(0x2D, 0x37, 0x48)

    def add_header(slide, title_text, category_text="GARANTIA E QUALIDADE DE SOFTWARE (GQS)"):
        # Header Box
        tb = slide.shapes.add_textbox(PPTInches(0.8), PPTInches(0.5), PPTInches(11.7), PPTInches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = PPTPt(12)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_BLUE

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = PPTPt(26)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE

    # SLIDE 1: Capa
    s1 = prs.slides.add_slide(blank_layout)
    s1.background.fill.solid()
    s1.background.fill.fore_color.rgb = BG_DARK

    tb1 = s1.shapes.add_textbox(PPTInches(1.0), PPTInches(1.2), PPTInches(11.3), PPTInches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CENTRO UNIVERSITÁRIO UNA — CAMPUS BARREIRO"
    p.font.size = PPTPt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p = tf1.add_paragraph()
    p.text = "EduGrade — Calculadora & Gestão Acadêmica"
    p.font.size = PPTPt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_after = PPTPt(14)

    p = tf1.add_paragraph()
    p.text = "Projeto Prático A3 | Disciplina: Garantia e Qualidade de Software (GQS)"
    p.font.size = PPTPt(18)
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = PPTPt(24)

    p = tf1.add_paragraph()
    p.text = "Professor Orientador: Daniel Henrique Matos de Paiva"
    p.font.size = PPTPt(14)
    p.font.color.rgb = TEXT_MUTED
    p.space_after = PPTPt(16)

    members_text = (
        "Integrantes da Equipe:\n"
        "• Lucas Henrique Miranda (RA: 325131396) — Product Owner & Requisitos\n"
        "• Caio Duraes (RA: 325132875) — Desenvolvedor Backend & Arquitetura\n"
        "• Gabriel Ferreira (RA: 325140970) — DevOps & Gestão GitFlow\n"
        "• Fernando Almeida (RA: 326132695) — Engenheiro de Qualidade & TDD"
    )
    p = tf1.add_paragraph()
    p.text = members_text
    p.font.size = PPTPt(13)
    p.font.color.rgb = TEXT_WHITE

    # SLIDE 2: Problema Real & ODS 4
    s2 = prs.slides.add_slide(blank_layout)
    s2.background.fill.solid()
    s2.background.fill.fore_color.rgb = BG_DARK
    add_header(s2, "Problema Real & Vínculo com ODS 4 (ONU)")

    tb2 = s2.shapes.add_textbox(PPTInches(0.8), PPTInches(1.8), PPTInches(11.7), PPTInches(5.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "O Desafio do Estudante Universitário:"
    p.font.size = PPTPt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = PPTPt(8)

    bullets = [
        "Sistemas de notas complexos com pesos variados (A1, A2, A3) geram falta de visibilidade do desempenho.",
        "Estudantes descobrem situação de risco apenas no final do semestre, impossibilitando a recuperação a tempo.",
        "Consequências graves: ansiedade acadêmica, reprovações desnecessárias e aumento da taxa de evasão."
    ]
    for b in bullets:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = PPTPt(15)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(6)

    p = tf2.add_paragraph()
    p.text = "\nAlinhamento com a ODS 4 (Educação de Qualidade):"
    p.font.size = PPTPt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = PPTPt(8)

    ods_bullets = [
        "Meta 4.4: Desenvolvimento de competências técnicas e autonomia no planejamento acadêmico.",
        "Incentivo à Permanência: Prevenção da evasão universitária mediante feedback preventivo e cálculo preditivo."
    ]
    for b in ods_bullets:
        p = tf2.add_paragraph()
        p.text = "• " + b
        p.font.size = PPTPt(15)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(6)

    # SLIDE 3: A Solução EduGrade
    s3 = prs.slides.add_slide(blank_layout)
    s3.background.fill.solid()
    s3.background.fill.fore_color.rgb = BG_DARK
    add_header(s3, "A Solução EduGrade — Funcionalidades & Arquitetura")

    tb3 = s3.shapes.add_textbox(PPTInches(0.8), PPTInches(1.8), PPTInches(11.7), PPTInches(5.0))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    sol_items = [
        ("Cálculo Ponderado & Personalizado:", "Cálculo preciso considerando pesos institucionais de cada avaliação."),
        ("Diagnóstico de Situação Automático:", "Classificação imediata em Aprovado, Recuperação ou Reprovado."),
        ("Cálculo Preditivo para Exame Final:", "Determina a nota mínima decimal necessária para aprovação na prova final."),
        ("Persistência em JSON:", "Armazenamento confiável do histórico de estudantes e disciplinas."),
        ("Arquitetura Modular Clean Code:", "Separação clara de responsabilidades em src/calculator.py, src/manager.py e src/exceptions.py.")
    ]
    for title, desc in sol_items:
        p = tf3.add_paragraph() if tf3.paragraphs[0].text else tf3.paragraphs[0]
        p.text = f"• {title} "
        p.font.size = PPTPt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(10)

    # SLIDE 4: GitFlow e Commits
    s4 = prs.slides.add_slide(blank_layout)
    s4.background.fill.solid()
    s4.background.fill.fore_color.rgb = BG_DARK
    add_header(s4, "Gestão de Configuração & GitFlow no GitHub")

    tb4 = s4.shapes.add_textbox(PPTInches(0.8), PPTInches(1.8), PPTInches(11.7), PPTInches(5.0))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    git_items = [
        ("Estrutura de Branches GitFlow:", "Utilização das branches main (produção), develop (desenvolvimento) e branches de feature."),
        ("Commits Semânticos:", "Histórico limpo e rastreável utilizando os prefixos convencionais feat:, fix:, test:, docs:."),
        ("Colaboração em Equipe:", "Divisão de tarefas e integração contínua via GitHub entre os 4 integrantes."),
        ("Repositório Oficial:", "https://github.com/vivoeasy100/calculadora_de_notas_academicas-")
    ]
    for title, desc in git_items:
        p = tf4.add_paragraph() if tf4.paragraphs[0].text else tf4.paragraphs[0]
        p.text = f"• {title} "
        p.font.size = PPTPt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(12)

    # SLIDE 5: Qualidade, TDD e CI/CD
    s5 = prs.slides.add_slide(blank_layout)
    s5.background.fill.solid()
    s5.background.fill.fore_color.rgb = BG_DARK
    add_header(s5, "Garantia da Qualidade, TDD e Automação CI/CD")

    tb5 = s5.shapes.add_textbox(PPTInches(0.8), PPTInches(1.8), PPTInches(11.7), PPTInches(5.0))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    qa_items = [
        ("Desenvolvimento Guiado por Testes (TDD):", "Criação prévia de suíte de testes com unittest cobrindo 100% dos cenários de cálculo."),
        ("Suíte de Testes Automatizados:", "10 testes executados em milissegundos validando casos limite, exceções e entradas inválidas."),
        ("Tratamento de Exceções Personalizado:", "Classes AcademicError, InvalidGradeError e InvalidWeightError impedem falhas inesperadas."),
        ("Pipeline de CI/CD (GitHub Actions):", "Execução automática de testes a cada push e pull request via .github/workflows/ci.yml.")
    ]
    for title, desc in qa_items:
        p = tf5.add_paragraph() if tf5.paragraphs[0].text else tf5.paragraphs[0]
        p.text = f"• {title} "
        p.font.size = PPTPt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(12)

    # SLIDE 6: Conclusão
    s6 = prs.slides.add_slide(blank_layout)
    s6.background.fill.solid()
    s6.background.fill.fore_color.rgb = BG_DARK
    add_header(s6, "Conclusão & Checklist das Entregas")

    tb6 = s6.shapes.add_textbox(PPTInches(0.8), PPTInches(1.8), PPTInches(11.7), PPTInches(5.0))
    tf6 = tb6.text_frame
    tf6.word_wrap = True

    chk_items = [
        ("✔ Relatório Técnico Word (.docx):", "Pronto e formatado com todas as seções e fundamentação teórica."),
        ("✔ Código Fonte & GitHub Versionado:", "Código limpo, estruturado em GitFlow e com CI via GitHub Actions."),
        ("✔ Apresentação em Slides (.pptx):", "Slides estruturados para a apresentação do grupo."),
        ("✔ Vídeo Pitch de 5 Minutos:", "Roteiro gravado e publicado no YouTube/Drive segundo a divisão de tempo.")
    ]
    for title, desc in chk_items:
        p = tf6.add_paragraph() if tf6.paragraphs[0].text else tf6.paragraphs[0]
        p.text = f"{title} "
        p.font.size = PPTPt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_WHITE
        p.space_after = PPTPt(14)

    p_end = tf6.add_paragraph()
    p_end.text = "\nObrigado! Estamos abertos a dúvidas e considerações."
    p_end.font.size = PPTPt(20)
    p_end.font.bold = True
    p_end.font.color.rgb = ACCENT_GREEN
    p_end.alignment = PP_ALIGN.CENTER

    prs.save("Apresentacao_A3_EduGrade.pptx")
    print("Apresentacao_A3_EduGrade.pptx criada com sucesso!")

if __name__ == "__main__":
    generate_word_document()
    generate_pptx_presentation()
